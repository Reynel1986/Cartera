from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

# Texto extraído y formateado
texto = [
    ('Prueba Final Español - Laboraliza', True),
    ('Nombre ___________________________  Grupo _____', False),
    ('I. Lee detenidamente.', True),
    ('Whitman demostró la poesía al celebrar lo cotidiano: el cuerpo, la naturaleza, la unidad y la armonía. Su lenguaje era directo, coloquial y la osadía provocaba lo que generó controversia en su época, su poesía buscaba contar lo espiritual con lo terrenal, planteando: "cada átomo mío es tuyo".', False),
    ('1.1 Marca V ( ) o F ( ) Justifique si falso a las siguientes afirmaciones:', False),
    ('a) ___ Walt Whitman utilizó principalmente la rima consonante y motivos clásicos.', False),
    ('b) ___ Su lenguaje era rebuscado y provocador.', False),
    ('c) ___ Conectaba con el espíritu.', False),
    ('', False),
    ('III. Del texto inicial extrae:', True),
    ('a) Dos pronombres', False),
    ('b) Una preposición', False),
    ('c) Analiza sintácticamente la siguiente oración:', False),
    ('Para todos su poesía representa a la colectividad de los seres humanos.', False),
    ('', False),
    ('IV. Dictado', True)
]

# Crear documento nuevo
doc = Document()

# Insertar el texto con formato
for linea, es_titulo in texto:
    p = doc.add_paragraph()
    run = p.add_run(linea)
    if es_titulo:
        run.bold = True
        run.font.size = Pt(14)
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    else:
        run.font.size = Pt(12)
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

doc.save('img/español-1.docx')
print('¡Documento actualizado con éxito!') 